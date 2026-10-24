import openpyxl
from collections import Counter

wb = openpyxl.load_workbook('output/master_karyawan.xlsx', data_only=True)
v = wb['Kumpulan Video']
links = [v.cell(i, 6).value for i in range(2, v.max_row + 1) if v.cell(i, 6).value]
counts = Counter(links)
dups = {k: c for k, c in counts.items() if c > 1}

print(f"Total baris: {len(links)}")
print(f"Link unik: {len(set(links))}")
print(f"Duplikat: {len(links) - len(set(links))}")
if dups:
    print("\nLink yang duplikat:")
    for link, c in dups.items():
        print(f"  {c}x: {link}")
else:
    print("Tidak ada duplikat.")
