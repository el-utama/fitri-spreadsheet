"""Perbaiki live master setelah migrasi kolom K.

1) MASTER: tulis ulang formula P2:S1000 dari master lokal.
2) RINGKASAN: pasang formula minimal C:J (Link Video Kosong / Bukan YouTube).
3) Header/note RINGKASAN dan helper status.

Prasyarat: MASTER live sudah 19 kolom (A:S) dengan K = Link Video YT.
"""
import json, sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl, pyperclip
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
FID = json.loads((ROOT / "data" / "drive_upload.json").read_text(encoding="utf-8"))["fileId"]
STATUS_LIST = json.loads((ROOT / "config" / "validation_lists.json").read_text(encoding="utf-8"))["status"]
STATUS_LAST = len(STATUS_LIST) + 1

PC_SHEET = {1: "PC 1", 2: "PC 02", 3: "PC 03", 4: "PC 04", 5: "PC 05", 6: "PC 06",
            7: "PC 07", 8: "PC 08", 9: "PC 09", 10: "PC 10", 11: "PC 11", 12: "PC 12"}
PC_IDS = {int(m["n"]): m["id"] for m in
          json.loads((ROOT / "data" / "manifest.json").read_text(encoding="utf-8"))}
R = 1000


def imp(i, rng):
    url = f"https://docs.google.com/spreadsheets/d/{PC_IDS[i]}/edit"
    return f'IMPORTRANGE("{url}","\'{PC_SHEET[i]}\'!{rng}")'


def ringkasan_rows():
    rows = []
    for i in range(1, 13):
        A, D, F = imp(i, f"A2:A{R}"), imp(i, f"D2:D{R}"), imp(i, f"F2:F{R}")
        I, J, K, L = (imp(i, f"I2:I{R}"), imp(i, f"J2:J{R}"),
                      imp(i, f"K2:K{R}"), imp(i, f"L2:L{R}"))
        rows.append([
            f'=SUMPRODUCT(--({D}<>""))',
            f'=SUMPRODUCT(({D}<>"")*({I}=""))',
            f'=SUMPRODUCT(({D}<>"")*({F}=""))',
            f'=SUMPRODUCT(({D}<>"")*({L}=""))',
            f'=SUMPRODUCT(({D}<>"")*({L}<>""))-SUMPRODUCT(COUNTIF({L},$N$2:$N${STATUS_LAST}))',
            f'=SUMPRODUCT(--(COUNTIF({A},{A})>1)*({A}<>""))',
            f'=SUMPRODUCT(({D}<>"")*({I}<>""))-SUMPRODUCT(--ISNUMBER(SEARCH("youtu",{I})))',
            f'=IFERROR(AVERAGE(FILTER({K}-{J},ISNUMBER({J}),ISNUMBER({K}),{K}>={J})),"")',
        ])
    return rows


paste_block = "\n".join("\t".join(r) for r in ringkasan_rows())

wb = openpyxl.load_workbook(ROOT / "output" / "master_karyawan.xlsx")
m = wb["MASTER"]
header = [m.cell(1, c).value for c in range(1, 20)]
if header[10] != "Link Video YT":
    raise SystemExit("Master lokal belum punya K = Link Video YT; jalankan build_master.py dulu.")
formulas = [m.cell(2, c).value for c in range(16, 20)]
if not all(isinstance(f, str) and f.startswith("=") for f in formulas):
    raise SystemExit("Formula P2:S2 master lokal tidak lengkap.")
formula_tsv = "\t".join(formulas)
print("formula P2:S2 siap, panjang:", [len(x) for x in formulas])

with sync_playwright() as pw:
    b = pw.chromium.connect_over_cdp("http://127.0.0.1:9222")
    ctx = b.contexts[0]
    p = ctx.new_page()
    p.goto(f"https://docs.google.com/spreadsheets/d/{FID}/edit",
           wait_until="domcontentloaded", timeout=90000)
    p.wait_for_timeout(6000)

    # ---- MASTER: P2:S1000 ----
    for t in p.query_selector_all(".docs-sheet-tab"):
        if (t.inner_text() or "").strip().startswith("MASTER"):
            t.click()
            break
    p.wait_for_timeout(2500)
    nb = p.locator("#t-name-box").first
    nb.click(); p.wait_for_timeout(250)
    nb.fill("P2"); p.keyboard.press("Enter")
    p.wait_for_timeout(600)
    pyperclip.copy(formula_tsv)
    p.keyboard.press("Control+v")
    p.wait_for_timeout(1000)
    p.keyboard.press("Escape")
    nb.click(); p.wait_for_timeout(250)
    nb.fill("P2:S1000"); p.keyboard.press("Enter")
    p.wait_for_timeout(800)
    p.keyboard.press("Control+d")
    p.wait_for_timeout(5000)
    p.keyboard.press("Escape")
    print("[MASTER] formula P2:S1000 dipasang")

    # ---- RINGKASAN ----
    for t in p.query_selector_all(".docs-sheet-tab"):
        if "RINGKASAN" in (t.inner_text() or ""):
            t.click()
            break
    p.wait_for_timeout(2500)
    nb.click(); p.wait_for_timeout(200)
    nb.fill("C2"); p.keyboard.press("Enter")
    p.wait_for_timeout(500)
    pyperclip.copy(paste_block)
    p.keyboard.press("Control+v")
    p.wait_for_timeout(7000)
    p.keyboard.press("Escape")

    # header C1:J1
    nb.click(); p.wait_for_timeout(200)
    nb.fill("C1"); p.keyboard.press("Enter")
    p.wait_for_timeout(300)
    headers = "\t".join(["Baris Data", "Link Video Kosong", "Genre Kosong", "Status Kosong",
                         "Status Invalid", "Duplikat", "Link Video Bukan YouTube",
                         "Rata-rata Masa Produksi"])
    pyperclip.copy(headers)
    p.keyboard.press("Control+v")
    p.wait_for_timeout(1500)
    p.keyboard.press("Escape")

    # helper status N1:N9 (8 status)
    nb.click(); p.wait_for_timeout(200)
    nb.fill("N1"); p.keyboard.press("Enter")
    p.wait_for_timeout(300)
    helper = "\n".join(["Daftar Status"] + STATUS_LIST)
    pyperclip.copy(helper)
    p.keyboard.press("Control+v")
    p.wait_for_timeout(1500)
    p.keyboard.press("Escape")

    # catatan
    nb.click(); p.wait_for_timeout(200)
    nb.fill("A15"); p.keyboard.press("Enter")
    p.wait_for_timeout(300)
    pyperclip.copy("Catatan: RINGKASAN dihitung real-time dari 12 file PC via IMPORTRANGE. "
                   "Link Video Kosong/Bukan YouTube dihitung dari kolom I. "
                   "Kumpulan Video memakai LET+QUERY dari MASTER dan hanya memuat link video unik (kemunculan pertama).")
    p.keyboard.press("Control+v")
    p.wait_for_timeout(1500)
    p.keyboard.press("Escape")
    p.wait_for_timeout(4000)
    try:
        p.screenshot(path=str(ROOT / "output" / "ringkasan_v6.png"), timeout=60000)
    except Exception as exc:
        print("screenshot gagal (tidak fatal):", exc)
    print("[RINGKASAN] formula C2:J13 + header/helper/catatan dipasang")
    p.close(); b.close()
