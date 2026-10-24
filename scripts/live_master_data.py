"""Paste snapshot data gabungan (A:O) ke tab MASTER di Google Sheet live.

Prasyarat:
- Master live sudah punya kolom K = Link Video YT (19 kolom A:S).
- Chrome CDP aktif di port 9222 dengan akun yang punya akses edit.
"""
import json, sys
from datetime import datetime, date
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl, pyperclip
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
FID = json.loads((ROOT / "data" / "drive_upload.json").read_text(encoding="utf-8"))["fileId"]

wb = openpyxl.load_workbook(ROOT / "output" / "master_karyawan.xlsx", data_only=True)
ws = wb["MASTER"]

header = [ws.cell(1, c).value for c in range(1, 20)]
if header[10] != "Link Video YT":
    raise SystemExit(
        "Header K master lokal bukan 'Link Video YT'. "
        "Bangun ulang master dengan build_master.py versi baru."
    )
if len([h for h in header if h not in (None, "")]) < 19:
    raise SystemExit("Master lokal belum 19 kolom A:S. Bangun ulang dengan build_master.py.")

rows = []
r = 2
while True:
    vals = [ws.cell(r, c).value for c in range(1, 16)]
    if all(v is None or str(v).strip() == "" for v in vals):
        break
    out = []
    for idx, v in enumerate(vals, 1):
        if v is None:
            out.append("")
        elif isinstance(v, (datetime, date)):
            if idx == 4 and isinstance(v, datetime):
                out.append(v.strftime("%Y-%m-%d %H:%M"))
            else:
                out.append(v.strftime("%Y-%m-%d"))
        else:
            out.append(str(v))
    rows.append(out)
    r += 1

print("baris untuk dipaste:", len(rows))
tsv = "\n".join("\t".join(rw) for rw in rows)
(ROOT / "output" / "master_live.tsv").write_text(tsv, encoding="utf-8")

with sync_playwright() as pw:
    b = pw.chromium.connect_over_cdp("http://127.0.0.1:9222")
    ctx = b.contexts[0]
    p = ctx.new_page()
    p.goto(f"https://docs.google.com/spreadsheets/d/{FID}/edit",
           wait_until="domcontentloaded", timeout=90000)
    p.wait_for_timeout(6000)
    for t in p.query_selector_all(".docs-sheet-tab"):
        if (t.inner_text() or "").strip().startswith("MASTER"):
            t.click()
            break
    p.wait_for_timeout(2500)
    nb = p.locator("#t-name-box").first
    nb.click(); p.wait_for_timeout(250)
    nb.fill("A2:O1000"); p.keyboard.press("Enter")
    p.wait_for_timeout(1200)
    p.keyboard.press("Delete")
    p.wait_for_timeout(1500)
    nb.click(); p.wait_for_timeout(250)
    nb.fill("A2"); p.keyboard.press("Enter")
    p.wait_for_timeout(800)
    pyperclip.copy(tsv)
    p.keyboard.press("Control+v")
    p.wait_for_timeout(7000)
    p.keyboard.press("Escape")
    p.wait_for_timeout(4000)
    p.screenshot(path=str(ROOT / "output" / "master_live_data.png"))
    print("[ok] data A:O dipaste ke MASTER live")
    p.close(); b.close()
