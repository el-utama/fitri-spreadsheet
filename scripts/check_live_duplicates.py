import io
import json
from collections import Counter
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parents[1]
FID = json.loads((ROOT / "data" / "drive_upload.json").read_text(encoding="utf-8"))["fileId"]
r = requests.get(f"https://docs.google.com/spreadsheets/d/{FID}/export?format=xlsx", timeout=90)
r.raise_for_status()
wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
ws = wb["Kumpulan Video"]
links = [ws.cell(i, 6).value for i in range(2, ws.max_row + 1) if ws.cell(i, 6).value]
counts = Counter(links)
dups = {k: c for k, c in counts.items() if c > 1}
print(f"Total baris data: {len(links)}")
print(f"Link unik: {len(set(links))}")
print(f"Duplikat: {len(links) - len(set(links))}")
if dups:
    for link, c in dups.items():
        print(f"  {c}x: {link}")
else:
    print("Tidak ada duplikat.")
