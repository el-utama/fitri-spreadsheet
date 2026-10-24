"""
Membangun spreadsheet master dari 12 sumber "PC 01..PC 12".

Tab master:
  PETUNJUK, DAFTAR_SUMBER, MASTER, Kumpulan Video, KAMUS_KOLOM, RINGKASAN,
  VALIDASI_TEMPLATE, TEMUAN_STRUKTUR, TEMPLET_PC, RAW_PC01..RAW_PC12
"""
import json
import sys
from datetime import datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
MANIFEST = json.loads((ROOT / "data" / "manifest.json").read_text(encoding="utf-8"))
OUTDIR = ROOT / "output"
OUTDIR.mkdir(exist_ok=True)
OUT = OUTDIR / "master_karyawan.xlsx"

PC_LABELS = [f"PC {i:02d}" for i in range(1, 13)]
LISTS = json.loads((ROOT / "config" / "validation_lists.json").read_text(encoding="utf-8"))
AGGREGATOR_LIST = LISTS["aggregator"]
GENRE_LIST = LISTS["genre"]
STATUS_LIST = LISTS["status"]
LAST = 1000

H_FILL = PatternFill("solid", start_color="1F77B4")
H_FONT = Font(bold=True, color="FFFFFF")
ZEBRA = PatternFill("solid", start_color="F2F7FB")
OK_FILL = PatternFill("solid", start_color="E8F5E9")
WARN_FILL = PatternFill("solid", start_color="FFF3E0")
SAMPLE_FILL = PatternFill("solid", start_color="FFF9C4")
THIN = Side(style="thin", color="B0BEC5")
BORD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_header(ws, labels, widths=None, notes=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=1, column=i, value=lab)
        c.font = H_FONT
        c.fill = H_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORD
        if notes and lab in notes:
            cm = Comment(notes[lab], "master-builder", height=120, width=300)
            c.comment = cm
    ws.row_dimensions[1].height = 32
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def style_body(ws, r0, r1, wrap_cols=()):
    for r in range(r0, r1 + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORD
            cell.alignment = Alignment(vertical="center", wrap_text=c in wrap_cols)
            if r % 2 == 1:
                if cell.fill.start_color.rgb in (None, "00000000") or cell.fill.patternType is None:
                    cell.fill = ZEBRA


def arr(items):
    return "{" + ",".join('"' + str(x).replace('"', '""') + '"' for x in items) + "}"


def genre_norm(cell):
    return (f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM({cell}),"/",","),";",",")," ","")')


def genre_count(cell):
    n = genre_norm(cell)
    return f'(LEN({n})-LEN(SUBSTITUTE({n},",",""))+1)'


def genre_search(cell):
    n = genre_norm(cell)
    return f'SUMPRODUCT(--ISNUMBER(SEARCH(","&{arr(GENRE_LIST)}&",",","&{n}&",")))'


def genre_dv_formula(cell="F2"):
    n = genre_norm(cell)
    c = genre_count(cell)
    s = genre_search(cell)
    return f'=OR(TRIM({cell})="",AND({c}<=2,{s}={c}))'


def genre_issue(cell="F2"):
    c = genre_count(cell)
    s = genre_search(cell)
    return (f'IF(TRIM({cell})="","Genre kosong",'
            f'IF({c}>2,"Genre maksimal 2",'
            f'IF({s}<>{c},"Genre di luar daftar","")))')


def aggregator_issue(cell="G2"):
    return (f'IF(TRIM({cell})="","",'
            f'IF(ISERROR(MATCH(TRIM({cell}),{arr(AGGREGATOR_LIST)},0)),'
            f'"Aggregator di luar daftar",""))')


def status_issue(cell="N2"):
    return (f'IF(TRIM({cell})="","Status kosong",'
            f'IF(ISERROR(MATCH(TRIM({cell}),{arr(STATUS_LIST)},0)),'
            f'"Status di luar daftar",""))')


def main():
    m_by_n = {m["n"]: m for m in MANIFEST}
    # data mentah per sumber (baris pertama = header asli, termasuk spasi ekor)
    raw_tabs = {}
    for m in MANIFEST:
        wb_src = openpyxl.load_workbook(RAW / f"src{m['n']:02d}.xlsx", data_only=True)
        s = wb_src.worksheets[0]
        raw_tabs[m["n"]] = [list(r) for r in s.iter_rows(values_only=True)]

    wb = openpyxl.Workbook()

    # ---------------- PETUNJUK ----------------
    ws = wb.active
    ws.title = "PETUNJUK"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 120
    PETUNJUK = [
        ("SPREADSHEET MASTER - KONSOLIDASI KARYAWAN (PC 01..PC 12)", 14, True),
        ("Versi template. 1 baris = 1 video/lagu. Kunci utama = PC + Worker ID.", 11, False),
        ("", 11, False),
        ("== 1. STRUKTUR FILE ==", 12, True),
        ("MASTER            : tab utama berisi gabungan semua PC; isian karyawan dilakukan di file masing-masing, master menyaring & merangkum.", 11, False),
        ("Kumpulan Video    : daftar baris yang punya Link Video YT; 9 kolom; hanya link video unik (kemunculan pertama).", 11, False),
        ("RAW_PC01..RAW_PC12: salinan data mentah tiap sumber (header asli dipertahankan, termasuk spasi ekor).", 11, False),
        ("RINGKASAN         : jumlah baris, kekurangan isi, duplikat, dan rata-rata masa produksi per PC.", 11, False),
        ("KAMUS_KOLOM       : pemetaan kolom sumber -> kolom master + tipe data + format + aturan isi.", 11, False),
        ("TEMUAN_STRUKTUR   : daftar masalah struktur/kolom yang ditemukan dan bagaimana ditangani di master.", 11, False),
        ("VALIDASI_TEMPLATE : kerangka validasi per PC (dicek Apps Script dan dijadikan tab VALIDASI di file karyawan).", 11, False),
        ("TEMPLET_PC        : template kolom baku untuk file PC (bila ingin disalin ke file karyawan).", 11, False),
        ("DAFTAR_SUMBER     : daftar 12 spreadsheet sumber dengan URL dan status.", 11, False),
        ("", 11, False),
        ("== 2. KOLOM MASTER DAN CARA ISI ==", 12, True),
        ("PC                    : dropdown PC 01..PC 12, wajib. Bagian dari kunci utama.", 11, False),
        ("Nama Karyawan         : teks, opsional. Boleh diisi kemudian; bermakna bila 1 karyawan menangani lebih dari 1 PC.", 11, False),
        ("Worker ID             : teks, wajib, unik di dalam PC. Format disarankan 'PC07-0001'.", 11, False),
        ("Update Date           : tanggal-jam, format YYYY-MM-DD HH:MM. Waktu entri terakhir diubah.", 11, False),
        ("Song ID YT            : ID video YouTube (bukan URL). Wajib.", 11, False),
        ("Judul Lagu Kompilasi  : teks judul, tanpa pindah baris. Wajib.", 11, False),
        ("Aggregator            : pilih dari daftar baku: " + ", ".join(AGGREGATOR_LIST) + ".", 11, False),
        ("Genre                 : pilih 1 atau 2 genre dari daftar baku, pisahkan dengan koma. Contoh: Pop, Jazz. Wajib.", 11, False),
        ("Nama Channel          : nama channel YouTube persis. Satu PC boleh mengelola >1 channel.", 11, False),
        ("Link Channel YT       : URL channel (bukan link video). Wajib. Boleh youtube.com atau youtu.be.", 11, False),
        ("Link Video YT         : URL video YouTube (bukan link channel). Wajib. Boleh youtube.com atau youtu.be.", 11, False),
        ("Tanggal Video Dibuat  : tanggal, format YYYY-MM-DD, wajib.", 11, False),
        ("Tanggal Video Dipublish: tanggal, format YYYY-MM-DD, boleh kosong sebelum tayang.", 11, False),
        ("Status Upload         : pilih dari daftar baku: " + ", ".join(STATUS_LIST) + ". Default Draft.", 11, False),
        ("PIC channel           : nama PIC channel, boleh kosong.", 11, False),
        ("Masa Produksi (hari)  : formula otomatis = Dipublish - Dibuat. Jangan diisi manual.", 11, False),
        ("Umur Data (hari)      : formula otomatis = hari ini - Update Date. Jangan diisi manual.", 11, False),
        ("Validasi Master       : formula otomatis, berisi daftar masalah baris; kosong jika bersih.", 11, False),
        ("Duplikat?             : formula otomatis, 'DUPLICAT' bila kombinasi (PC, Worker ID) dipakai lebih dari satu baris.", 11, False),
        ("", 11, False),
        ("== 3. VALIDASI OTOMATIS ==", 12, True),
        ("Baris sample kuning di MASTER hanyalah contoh; hapus/replace begitu data asli masuk.", 11, False),
        ("Kolom Validasi Master memeriksa: kolom wajib kosong (termasuk Update Date, Song ID, Tgl Dibuat, Status), link bukan youtube, tanggal tidak bertipe tanggal, publish < dibuat, aggregator/genre/status di luar daftar, dan genre lebih dari 2.", 11, False),
        ("Link channel dan link video sama-sama menerima youtube.com maupun youtu.be.", 11, False),
        ("File PC juga bisa memakai tab VALIDASI (lihat apps_script/karyawan_validasi.gs).", 11, False),
        ("", 11, False),
        ("== 4. CARA MEMPERBARUI MASTER SAAT DATA SUDAH DIISI ==", 12, True),
        ("python scripts/fetch_sheets.py   -> unduh ulang 12 file ke data/raw/", 11, False),
        ("python scripts/inventory.py      -> refresh struktur", 11, False),
        ("python scripts/build_master.py   -> bangun master + tab Kumpulan Video", 11, False),
        ("python scripts/merge_data.py     -> gabungkan isi baris data ke MASTER, RAW_PC*, dan Kumpulan Video", 11, False),
        ("Versi live memakai IMPORTRANGE untuk RINGKASAN; Kumpulan Video memakai LET+QUERY dari MASTER dan membuang link duplikat.", 11, False),
        ("", 11, False),
        ("== 5. APLIKASI KARYAWAN (APPS SCRIPT) ==", 12, True),
        ("File apps_script/karyawan_validasi.gs dipasang sekali di script.google.com.", 11, False),
        ("Ia melakukan: 1) catatan cara isi pada tiap header file PC; 2) tab VALIDASI per file; 3) komentar otomatis saat data tidak sesuai (onEdit).", 11, False),
    ]
    for i, (txt, size, bold) in enumerate(PETUNJUK, 1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font = Font(size=size, bold=bold)
        c.alignment = Alignment(vertical="center", wrap_text=False)

    # ---------------- DAFTAR_SUMBER ----------------
    ws = wb.create_sheet("DAFTAR_SUMBER")
    write_header(ws, ["PC", "Spreadsheet ID", "URL", "Baris Data", "Kolom"],
                 widths=[10, 48, 80, 12, 8])
    for i, m in enumerate(MANIFEST, start=2):
        rows = raw_tabs[m["n"]]
        # hitung baris data sebenarnya (punya Worker ID / Song ID / Judul),
        # bukan baris berformat kosong atau 975 baris sampah PC 09.
        n_rows = sum(
            1 for row in rows[1:]
            if any((row[idx] if idx < len(row) else None) not in (None, "")
                   for idx in (0, 2, 3))
        )
        n_cols = sum(1 for h in (rows[0] if rows else []) if h not in (None, ""))
        vals = [PC_LABELS[i - 2], m["id"], f"https://docs.google.com/spreadsheets/d/{m['id']}/edit",
                n_rows, n_cols]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORD
        link = ws.cell(row=i, column=3)
        link.hyperlink = vals[2]
        link.font = Font(color="1F77B4", underline="single", size=11)

    # ---------------- KAMUS_KOLOM ----------------
    ws = wb.create_sheet("KAMUS_KOLOM")
    kamus = [
        ("PC", "nama file/tab sumber", "enum", "PC NN (01..12)", "wajib", "dropdown; tidak boleh di luar daftar"),
        ("Nama Karyawan", "baru (opsional)", "teks", "nama lengkap", "opsional", "satu karyawan boleh beberapa PC"),
        ("Worker ID", "Worker ID sumber", "teks", "PC07-0001 (disarankan)", "wajib", "unik di dalam PC"),
        ("Update Date", "Update Date sumber", "tanggal-jam", "YYYY-MM-DD HH:MM", "wajib", "harus benar-benar bertipe tanggal"),
        ("Song ID YT", "Song ID YT sumber (trim spasi)", "teks", "ID video, contoh x7pB0nE1g4M", "wajib", "bukan URL"),
        ("Judul Lagu Kompilasi", "Judul Lagu Kompilasi sumber", "teks", "judul bebas", "wajib", "tanpa pindah baris"),
        ("Aggregator", "Aggregator sumber", "enum", ", ".join(AGGREGATOR_LIST), "opsional", "dropdown daftar baku"),
        ("Genre", "Genre sumber (trim spasi)", "enum", ", ".join(GENRE_LIST), "wajib", "pilih 1-2, pisahkan koma"),
        ("Nama Channel", "Nama Channel sumber", "teks", "nama persis channel", "wajib", "satu PC bisa >1 channel"),
        ("Link Channel YT", "Link Channel YT sumber", "URL", "https://www.youtube.com/@... atau https://youtu.be/...", "wajib", "boleh youtube.com atau youtu.be"),
        ("Link Video YT", "Link Video YT sumber", "URL", "https://www.youtube.com/watch?v=... atau https://youtu.be/...", "wajib", "boleh youtube.com atau youtu.be"),
        ("Tanggal Video Dibuat", "Tanggal Video Dibuat sumber", "tanggal", "YYYY-MM-DD", "wajib", "harus bertipe tanggal"),
        ("Tanggal Video Dipublish", "Tanggal Video Dipublish sumber", "tanggal", "YYYY-MM-DD", "opsional", ">= tanggal Dibuat"),
        ("Status Upload", "Status Upload sumber", "enum", ", ".join(STATUS_LIST), "wajib", "dropdown; default Draft"),
        ("PIC channel", "PIC channel sumber (trim spasi)", "teks", "nama PIC", "opsional", "boleh kosong"),
        ("Masa Produksi (hari)", "formula P", "angka", "hari", "otomatis", "Dipublish - Dibuat"),
        ("Umur Data (hari)", "formula Q", "angka", "hari", "otomatis", "hari ini - Update Date"),
        ("Validasi Master", "formula R", "teks", "daftar masalah", "otomatis", "termasuk Update Date, Song ID, Tgl Dibuat, Status; kosong bila bersih"),
        ("Duplikat?", "formula S", "enum", "DUPLICAT / kosong", "otomatis", "berdasarkan (PC, Worker ID)"),
    ]
    write_header(ws, ["Kolom Master", "Asal Kolom", "Tipe", "Format/Nilai", "Wajib?", "Aturan"],
                 widths=[26, 34, 14, 40, 12, 52])
    for i, row in enumerate(kamus, start=2):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORD
            c.alignment = Alignment(vertical="center", wrap_text=True)
        wajib = str(row[4]).lower()
        ws.cell(row=i, column=5).fill = OK_FILL if wajib == "wajib" else WARN_FILL
    ws.auto_filter.ref = f"A1:F{1+len(kamus)}"

    # ---------------- MASTER ----------------
    ws = wb.create_sheet("MASTER")
    mcols = ["PC", "Nama Karyawan", "Worker ID", "Update Date", "Song ID YT", "Judul Lagu Kompilasi",
             "Aggregator", "Genre", "Nama Channel", "Link Channel YT", "Link Video YT",
             "Tanggal Video Dibuat", "Tanggal Video Dipublish", "Status Upload", "PIC channel",
             "Masa Produksi (hari)", "Umur Data (hari)", "Validasi Master", "Duplikat?"]
    mwidths = [10, 20, 14, 19, 16, 34, 14, 12, 22, 36, 36, 15, 16, 16, 18, 12, 10, 42, 12]
    mnotes = {
        "PC": "Dropdown: PC 01..PC 12. Wajib. Bagian kunci utama (PC + Worker ID).",
        "Worker ID": "Harus unik di dalam PC. Format disarankan 'PC07-0001'.",
        "Update Date": "Format: YYYY-MM-DD HH:MM. Harus bertipe tanggal, bukan teks.",
        "Song ID YT": "ID video YouTube (bukan URL).",
        "Judul Lagu Kompilasi": "Judul; jangan pindah baris di dalam sel.",
        "Aggregator": "Dropdown daftar baku: " + ", ".join(AGGREGATOR_LIST) + ".",
        "Genre": "Pilih 1 atau 2 genre dari daftar baku; pisahkan dengan koma. Contoh: Pop, Jazz.",
        "Nama Channel": "Nama channel persis yang dipakai di baris ini.",
        "Link Channel YT": "URL channel (bukan link video). Wajib. Boleh youtube.com atau youtu.be.",
        "Link Video YT": "URL video YouTube (bukan link channel). Wajib. Boleh youtube.com atau youtu.be.",
        "Tanggal Video Dibuat": "Format: YYYY-MM-DD. Harus bertipe tanggal.",
        "Tanggal Video Dipublish": "Format: YYYY-MM-DD. Boleh kosong sebelum tayang.",
        "Status Upload": "Dropdown: " + ", ".join(STATUS_LIST),
        "PIC channel": "Nama PIC channel; boleh kosong.",
        "Masa Produksi (hari)": "FORMULA OTOMATIS kolom P. Jangan diisi manual.",
        "Umur Data (hari)": "FORMULA OTOMATIS kolom Q. Jangan diisi manual.",
        "Validasi Master": "FORMULA OTOMATIS kolom R berisi daftar masalah baris, termasuk Update Date/Song ID/Tgl Dibuat/Status. Kosong = bersih.",
        "Duplikat?": "FORMULA OTOMATIS kolom S menandai (PC, Worker ID) yang berulang.",
    }
    write_header(ws, mcols, widths=mwidths, notes=mnotes)

    dv_pc = DataValidation(type="list", formula1='"' + ",".join(PC_LABELS) + '"',
                           allow_blank=False, showErrorMessage=True,
                           errorTitle="PC tidak valid",
                           error="Pilih salah satu: " + ", ".join(PC_LABELS))
    ws.add_data_validation(dv_pc)
    dv_pc.add(f"A2:A{LAST}")
    dv_status = DataValidation(type="list", formula1='"' + ",".join(STATUS_LIST) + '"',
                               allow_blank=True, showErrorMessage=True,
                               errorTitle="Status tidak valid",
                               error="Pilih dari: " + ", ".join(STATUS_LIST))
    ws.add_data_validation(dv_status)
    dv_status.add(f"N2:N{LAST}")
    dv_agg = DataValidation(type="list", formula1='"' + ",".join(AGGREGATOR_LIST) + '"',
                            allow_blank=True, showErrorMessage=True,
                            errorTitle="Aggregator tidak valid",
                            error="Pilih dari: " + ", ".join(AGGREGATOR_LIST))
    ws.add_data_validation(dv_agg)
    dv_agg.add(f"G2:G{LAST}")
    dv_genre = DataValidation(type="custom", formula1=genre_dv_formula("H2"),
                              allow_blank=True, showErrorMessage=True,
                              errorTitle="Genre tidak valid",
                              error="Pilih 1-2 genre dari daftar; pisahkan dengan koma.")
    ws.add_data_validation(dv_genre)
    dv_genre.add(f"H2:H{LAST}")
    status_arr = arr(STATUS_LIST)

    for r in range(2, LAST + 1):
        ws.cell(row=r, column=4).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=r, column=12).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=13).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=16).value = (
            f'=IF(AND(ISNUMBER($L{r}),ISNUMBER($M{r})),MAX(0,$M{r}-$L{r}),"")')
        ws.cell(row=r, column=16).number_format = "0"
        ws.cell(row=r, column=17).value = (
            f'=IF(ISNUMBER($D{r}),MAX(0,TODAY()-$D{r}),"")')
        ws.cell(row=r, column=17).number_format = "0"
        ws.cell(row=r, column=18).value = (
            f'=IF(COUNTA($A{r}:$O{r})=0,"",IFERROR(TRIM(TEXTJOIN(" | ",TRUE,'
            f'IF($A{r}="","PC kosong",""),'
            f'IF($C{r}="","Worker ID kosong",""),'
            f'IF($D{r}="","Update Date kosong",IF(NOT(ISNUMBER($D{r})),"Update Date bukan tanggal","")),'
            f'IF($E{r}="","Song ID YT kosong",""),'
            f'IF($F{r}="","Judul kosong",""),'
            f'{aggregator_issue(f"$G{r}")},'
            f'{genre_issue(f"$H{r}")},'
            f'IF($I{r}="","Nama Channel kosong",""),'
            f'IF($J{r}="","Link Channel kosong",IF(ISERROR(SEARCH("youtu",$J{r})),"Link Channel bukan youtube","")),'
            f'IF($K{r}="","Link Video kosong",IF(ISERROR(SEARCH("youtu",$K{r})),"Link Video bukan youtube","")),'
            f'IF($L{r}="","Tgl Dibuat kosong",IF(NOT(ISNUMBER($L{r})),"Tgl Dibuat bukan tanggal","")),'
            f'IF(AND($M{r}<>"",NOT(ISNUMBER($M{r}))),"Tgl Publish bukan tanggal",""),'
            f'IF(AND(ISNUMBER($L{r}),ISNUMBER($M{r}),$M{r}<$L{r}),"Tgl Publish < Tgl buat",""),'
            f'{status_issue(f"$N{r}")}'
            f')),""))')
        ws.cell(row=r, column=19).value = (
            f'=IF(AND($A{r}<>"",$C{r}<>""),IF(COUNTIFS($A$2:$A${LAST},$A{r},$C$2:$C${LAST},$C{r})>1,"DUPLICAT",""),"")')
        ws.cell(row=r, column=18).alignment = Alignment(wrap_text=True, vertical="top")

    # baris sample (kuning)
    sample = ["PC 07", "", "PC07-0001", datetime(2026, 10, 22, 9, 0), "x7pB0nE1g4M",
              "CONTOH BARIS - hapus saat mulai isi", "TuneCore", "Pop",
              "Nama Channel Contoh", "https://www.youtube.com/@namachannel",
              "https://youtu.be/contoh123", datetime(2026, 10, 1), datetime(2026, 10, 5),
              "Draft", ""]
    for j, v in enumerate(sample, 1):
        ws.cell(row=2, column=j, value=v)
    for j in range(1, 20):
        ws.cell(row=2, column=j).fill = SAMPLE_FILL
    style_body(ws, 2, LAST, wrap_cols=(6, 10, 11, 18))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:S{LAST+1}"

    # ---------------- Kumpulan Video ----------------
    ws = wb.create_sheet("Kumpulan Video")
    vcols = ["PC", "Worker ID", "Song ID YT", "Judul Lagu Kompilasi", "Nama Channel",
             "Link Video YT", "Status Upload", "Tanggal Video Dibuat", "Tanggal Video Dipublish"]
    vwidths = [10, 14, 16, 34, 22, 36, 16, 18, 20]
    vnotes = {
        "Link Video YT": "Hanya baris dengan Link Video YT terisi dan unik; kemunculan pertama dipertahankan. Diisi otomatis oleh merge_data.py.",
        "Tanggal Video Dibuat": "Format: YYYY-MM-DD.",
        "Tanggal Video Dipublish": "Format: YYYY-MM-DD.",
    }
    write_header(ws, vcols, widths=vwidths, notes=vnotes)
    for r in range(2, LAST + 1):
        ws.cell(row=r, column=8).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=9).number_format = "yyyy-mm-dd"
    style_body(ws, 2, LAST, wrap_cols=(4, 5, 6))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{LAST+1}"

    # ---------------- RINGKASAN ----------------
    ws = wb.create_sheet("RINGKASAN")
    write_header(ws, ["PC", "Spreadsheet ID", "Baris Data", "Link Video Kosong", "Genre Kosong",
                      "Status Kosong", "Status Invalid", "Duplikat", "Link Video Bukan YouTube",
                      "Rata-rata Masa Produksi"],
                 widths=[10, 48, 12, 16, 14, 14, 14, 12, 22, 20])
    for i in range(12):
        r = i + 2
        raw = f"RAW_PC{i+1:02d}"
        vals = [
            PC_LABELS[i],
            m_by_n[i + 1]["id"],
            f"=SUMPRODUCT(--('{raw}'!$D$2:$D${LAST}<>\"\"))",
            f"=SUMPRODUCT(('{raw}'!$D$2:$D${LAST}<>\"\")*('{raw}'!$I$2:$I${LAST}=\"\"))",
            f"=SUMPRODUCT(('{raw}'!$D$2:$D${LAST}<>\"\")*('{raw}'!$F$2:$F${LAST}=\"\"))",
            f"=SUMPRODUCT(('{raw}'!$D$2:$D${LAST}<>\"\")*('{raw}'!$L$2:$L${LAST}=\"\"))",
            f"=SUMPRODUCT(('{raw}'!$D$2:$D${LAST}<>\"\")*('{raw}'!$L$2:$L${LAST}<>\"\")"
            f"*ISERROR(MATCH('{raw}'!$L$2:$L${LAST},{status_arr},0)))",
            f"=SUMPRODUCT(--(COUNTIF('{raw}'!$A$2:$A${LAST},'{raw}'!$A$2:$A${LAST})>1)"
            f"*('{raw}'!$A$2:$A${LAST}<>\"\"))",
            f"=SUMPRODUCT(('{raw}'!$D$2:$D${LAST}<>\"\")*('{raw}'!$I$2:$I${LAST}<>\"\"))"
            f"-SUMPRODUCT(--ISNUMBER(SEARCH(\"youtu\",'{raw}'!$I$2:$I${LAST}&\"\")))",
            f"=IFERROR(SUMPRODUCT(ISNUMBER('{raw}'!$J$2:$J${LAST})*ISNUMBER('{raw}'!$K$2:$K${LAST})"
            f"*('{raw}'!$K$2:$K${LAST}>='{raw}'!$J$2:$J${LAST})"
            f"*('{raw}'!$K$2:$K${LAST}-'{raw}'!$J$2:$J${LAST}))"
            f"/SUMPRODUCT(ISNUMBER('{raw}'!$J$2:$J${LAST})*ISNUMBER('{raw}'!$K$2:$K${LAST})"
            f"*('{raw}'!$K$2:$K${LAST}>='{raw}'!$J$2:$J${LAST})),\"\")",
        ]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = BORD
    note = ws.cell(row=15, column=1,
                   value="Catatan: RINGKASAN lokal dihitung dari RAW_PC01..RAW_PC12; versi live memakai IMPORTRANGE. "
                         "Link Video Kosong/Bukan YouTube dihitung dari kolom I. "
                         "Kumpulan Video memakai LET+QUERY dari MASTER dan hanya memuat link video unik (kemunculan pertama).")
    ws.merge_cells(start_row=15, start_column=1, end_row=15, end_column=10)
    note.alignment = Alignment(wrap_text=True)

    # ---------------- VALIDASI_TEMPLATE ----------------
    ws = wb.create_sheet("VALIDASI_TEMPLATE")
    write_header(ws, ["Keterangan"], widths=[120])
    ws.cell(row=2, column=1, value="Apps Script akan membuat tab VALIDASI di tiap file PC memakai blok formula di bawah.").alignment = Alignment(wrap_text=False)
    ws.cell(row=3, column=1, value="Formula 'Masalah Ditemukan' per baris (tarik ke bawah):").font = Font(bold=True)
    ftxt = ('=IF(COUNTA($A2:$M2)<=1,"",TEXTJOIN(" | ",TRUE,'
            'IF($A2="","Worker ID kosong",""),'
            'IF($B2="","Update Date kosong",IF(NOT(ISNUMBER($B2)),"Update Date bukan tanggal","")),'
            'IF($C2="","Song ID YT kosong",""),'
            'IF($D2="","Judul kosong",""),'
            + aggregator_issue("$E2") + ','
            + genre_issue("$F2") + ','
            'IF($G2="","Nama Channel kosong",""),'
            'IF($H2="","Link Channel kosong",IF(ISERROR(SEARCH("youtu",$H2)),"Link Channel bukan youtube","")),'
            'IF($I2="","Link Video kosong",IF(ISERROR(SEARCH("youtu",$I2)),"Link Video bukan youtube","")),'
            'IF($J2="","Tgl Dibuat kosong",IF(NOT(ISNUMBER($J2)),"Tgl Dibuat bukan tanggal","")),'
            'IF($K2="","",IF(NOT(ISNUMBER($K2)),"Tgl Publish bukan tanggal","")),'
            'IF(AND(ISNUMBER($J2),ISNUMBER($K2),$K2<$J2),"Tgl Publish < Tgl buat",""),'
            + status_issue("$L2") + '))')
    ws.cell(row=4, column=1, value=ftxt)
    ws.cell(row=4, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 80

    # ---------------- TEMUAN_STRUKTUR ----------------
    ws = wb.create_sheet("TEMUAN_STRUKTUR")
    write_header(ws, ["No", "Temuan", "Dampak / Penyebab", "Kategori", "Perlakuan di Master"],
                 widths=[5, 52, 60, 16, 50])
    TEMUAN = [
        ("Spasi ekor pada header sumber (Song ID YT / Genre / PIC channel)", "Cek resep pembanding kolom; formula menonjok tidak cocok.", "Struktur",
         "Kolom master dibersihkan (tanpa spasi ekor); sumber tidak diubah."),
        ("Nama tab tidak konsisten (PC 1 vs PC 02)", "Formula antar file bisa patah bila nama tab berubah.", "Struktur",
         "Master memakai label PC 01..PC 12; source tetap dipertahankan."),
        ("Kolom 'Link YT' ambigu (channel / video)", "Berbahaya bagi analisis karena tipe data tidak eksplisit.", "Struktur",
         "Sumber sudah dipecah: H = Link Channel YT, I = Link Video YT; master J/K."),
        ("Note header I masih menyalin note Link Channel YT", "Petunjuk isi kolom video salah.", "Struktur",
         "installAll versi 13 kolom memperbaiki note I; note J-M ikut bergeser."),
        ("Link video PC 09 sempat hilang dari sumber", "Saat audit kolom I kosong; data lama hanya ada di backup.", "Data",
         "Saat verifikasi kolom I sudah terisi 24/24; backup output/backup/pc09_link_lama.csv tetap disimpan."),
        ("Tanggal berisiko disimpan sebagai teks", "Rumus selisih tanggal & pengurutan gagal.", "Data",
         "Master memakai number-format date + pemeriksaan Validasi Master."),
        ("Daftar status/aggregator/genre belum seragam", "Analisis agregat sulit karena variasi penulisan dan dropdown berbeda antar file.", "Struktur",
         "Daftar baku disimpan di config/validation_lists.json; dropdown master dan Apps Script memakai daftar yang sama."),
        ("Genre hanya mendukung satu nilai", "Kebutuhan genre silang (mis. Pop, Jazz) tidak bisa dicatat rapi.", "Struktur",
         "Validasi genre 1-2 nilai, dipisahkan koma; dropdown custom di master dan file PC."),
        ("Tab tersembunyi 'Extract' dan 'Refresh & manage'", "Berisikan header sama namun kosong; membingungkan parser.", "Struktur",
         "Tidak dimasukkan ke master; bisa diarsip melalui Apps Script jika diinginkan."),
        ("Sebagian file PC sudah berisi data", "PC 01 = 20 baris (16 link video); PC 09 = 24 baris (24 link video); PC lain header saja.", "Data",
         "Master tetap dibangun sebagai template; RINGKASAN live terisi setelah data masuk."),
        ("Tidak ada kolom Nama Karyawan di sumber", "Hubungan PC <-> karyawan tidak eksplisit.", "Struktur",
         "Master menambah kolom 'Nama Karyawan' (opsional)."),
        ("Worker ID berpotensi tidak unik bila digabung antar PC", "Baris bisa saling menindih saat merge.", "Struktur",
         "Kunci master = (PC, Worker ID); flag DUPLICAT otomatis."),
        ("Tidak ada kolom catatan/inspection note", "Detail kejadian (revisi dsb.) tidak tercatat.", "Struktur",
         "Master menyarankan kolom catatan di file PC; Apps Script menyediakan tarikan."),
        ("PIC channel vs Nama Karyawan tidak sepola", "Satu PC bisa multi channel; satu channel bisa multi PIC.", "Struktur",
         "Master menyediakan dua kolom terpisah (PC + PIC channel)."),
        ("Link video duplikat di PC 01", "7 kemunculan duplikat; hanya 9 link unik. Tanpa dedup, Kumpulan Video berisi link berulang.", "Data",
         "Kumpulan Video dedup exact URL dan mempertahankan kemunculan pertama; sumber tidak diubah."),
        ("19 baris PC 01 tanpa Worker ID", "Kunci (PC, Worker ID) tidak dapat dipakai; Duplikat? tidak menandai sel kosong.", "Data",
         "Validasi Master menandai Worker ID kosong; sumber perlu dilengkapi."),
        ("975 baris sampah PC 09 (hanya Aggregator Tuncore)", "Mengganggu getLastRow dan VALIDASI lama.", "Data",
         "Script 13 kolom mengabaikan baris tanpa Worker ID/Song ID/Judul; tab VALIDASI dibersihkan ulang."),
        ("Status Upload / Genre di PC 01 di luar daftar lama", "3 baris Upload dan 5 baris JAZZ tidak sesuai daftar lama.", "Data",
         "Daftar diperluas: Upload dan Jazz valid; pencocokan tidak peka huruf besar/kecil."),
        ("Manifest tab tidak akurat", "PC 02-12 tidak mencantumkan tab data; dapat menyesatkan proses inventaris.", "Struktur",
         "fetch_sheets.py diubah membaca sheetnames dari XLSX; manifest diperbarui."),
        ("12 trigger lama dimiliki akun owner (Srijapar)", "Akun editor tidak dapat menghapus trigger milik akun lain; trigger menargetkan 12 file.", "Operasional",
         "Kode project PC 01 diperbarui ke 13 kolom; trigger lama tetap dipakai karena sudah mencakup 12 file. Jika owner menghapusnya, pasang ulang per file."),
    ]
    for i, (t, d, k, s) in enumerate(TEMUAN, start=2):
        vals = [i - 1, t, d, k, s]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORD
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row=i, column=4).fill = WARN_FILL if k == "Struktur" else OK_FILL
    ws.auto_filter.ref = f"A1:E{1+len(TEMUAN)}"

    # ---------------- TEMPLET_PC ----------------
    ws = wb.create_sheet("TEMPLET_PC")
    tcols = ["Worker ID", "Update Date", "Song ID YT", "Judul Lagu Kompilasi", "Aggregator",
             "Genre", "Nama Channel", "Link Channel YT", "Link Video YT",
             "Tanggal Video Dibuat", "Tanggal Video Dipublish", "Status Upload", "PIC channel"]
    twidths = [14, 19, 16, 34, 14, 12, 22, 36, 36, 15, 16, 16, 18]
    tnotes = {
        "Worker ID": "Format disarankan 'PC07-0001' (prefix PC + nomor urut 4 digit).",
        "Update Date": "Format: YYYY-MM-DD HH:MM.",
        "Song ID YT": "ID video YouTube, bukan URL.",
        "Judul Lagu Kompilasi": "Judul; tanpa pindah baris dalam sel.",
        "Aggregator": "Dropdown daftar baku: " + ", ".join(AGGREGATOR_LIST) + ".",
        "Genre": "Pilih 1 atau 2 genre dari daftar baku; pisahkan dengan koma. Contoh: Pop, Jazz.",
        "Nama Channel": "Nama channel persis.",
        "Link Channel YT": "URL channel (bukan link video); boleh youtube.com atau youtu.be.",
        "Link Video YT": "URL video YouTube; wajib; boleh youtube.com atau youtu.be.",
        "Tanggal Video Dibuat": "Format: YYYY-MM-DD.",
        "Tanggal Video Dipublish": "Format: YYYY-MM-DD; kosong bila belum tayang.",
        "Status Upload": "Dropdown: " + ", ".join(STATUS_LIST),
        "PIC channel": "Nama PIC; boleh kosong.",
    }
    write_header(ws, tcols, widths=twidths, notes=tnotes)
    dv_a = DataValidation(type="list", formula1='"' + ",".join(AGGREGATOR_LIST) + '"',
                          allow_blank=True, showErrorMessage=True, errorTitle="Aggregator tidak valid",
                          error="Pilih dari: " + ", ".join(AGGREGATOR_LIST))
    ws.add_data_validation(dv_a)
    dv_a.add(f"E2:E{LAST}")
    dv_g = DataValidation(type="custom", formula1=genre_dv_formula("F2"),
                          allow_blank=True, showErrorMessage=True, errorTitle="Genre tidak valid",
                          error="Pilih 1-2 genre dari daftar; pisahkan dengan koma.")
    ws.add_data_validation(dv_g)
    dv_g.add(f"F2:F{LAST}")
    dv_s2 = DataValidation(type="list", formula1='"' + ",".join(STATUS_LIST) + '"',
                           allow_blank=True, showErrorMessage=True,
                           errorTitle="Status tidak valid", error="Pilih dari: " + ", ".join(STATUS_LIST))
    ws.add_data_validation(dv_s2)
    dv_s2.add(f"L2:L{LAST}")
    dv_h = DataValidation(type="custom", formula1='=OR(LEFT(H2,4)="http",H2="")',
                          allow_blank=True, showErrorMessage=True,
                          errorTitle="Link tidak valid", error="Harus mulai dengan http.")
    ws.add_data_validation(dv_h)
    dv_h.add(f"H2:H{LAST}")
    dv_i = DataValidation(type="custom", formula1='=OR(LEFT(I2,4)="http",I2="")',
                          allow_blank=True, showErrorMessage=True,
                          errorTitle="Link tidak valid", error="Harus mulai dengan http.")
    ws.add_data_validation(dv_i)
    dv_i.add(f"I2:I{LAST}")
    for r in range(2, LAST + 1):
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=r, column=10).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=11).number_format = "yyyy-mm-dd"
    style_body(ws, 2, LAST, wrap_cols=(4, 8, 9))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:M{LAST+1}"

    # ---------------- RAW_PCxx ----------------
    for i in range(1, 13):
        rows = raw_tabs[i]
        ws = wb.create_sheet(f"RAW_PC{i:02d}")
        cols = [str(h) if h not in (None, "") else f"col{j}" for j, h in enumerate(rows[0], 1)] if rows else ["col1"]
        write_header(ws, cols, widths=[16] * max(1, len(cols)))
        for r, row in enumerate(rows[1:], 2):
            for j, v in enumerate(row, 1):
                if v not in (None, ""):
                    ws.cell(row=r, column=j, value=v)
        ws.sheet_properties.tabColor = "8E24AA"

    order = ["PETUNJUK", "DAFTAR_SUMBER", "MASTER", "Kumpulan Video", "KAMUS_KOLOM", "RINGKASAN",
             "VALIDASI_TEMPLATE", "TEMUAN_STRUKTUR", "TEMPLET_PC"] + [f"RAW_PC{i:02d}" for i in range(1, 13)]
    wb._sheets = [wb[t] for t in order]
    wb.active = wb["PETUNJUK"]
    colors = {"PETUNJUK": "1F77B4", "DAFTAR_SUMBER": "1F77B4", "MASTER": "43A047",
              "Kumpulan Video": "D81B60", "KAMUS_KOLOM": "FB8C00", "RINGKASAN": "5E35B1",
              "VALIDASI_TEMPLATE": "FB8C00", "TEMUAN_STRUKTUR": "E53935", "TEMPLET_PC": "00897B"}
    for t, col in colors.items():
        wb[t].sheet_properties.tabColor = col
    wb.save(OUT)
    print("saved ->", OUT)


if __name__ == "__main__":
    main()
