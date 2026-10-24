"""Verifikasi hasil installAll di 12 file PC (skema 13 kolom).

Cek:
- tab VALIDASI ada/tidak
- isi tab VALIDASI (baris bermasalah)
- catatan (note) di header & sel data, termasuk note I = Link Video
- warna merah muda pada sel bermasalah
"""
import io
import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = json.loads((ROOT / "data" / "manifest.json").read_text(encoding="utf-8"))
IDS = {int(m["n"]): m["id"] for m in MANIFEST}

NON_DATA = {"Extract", "Refresh & manage", "VALIDASI"}


def is_red(cell):
    try:
        f = cell.fill
        if not f or not f.fgColor:
            return False
        rgb = (f.fgColor.rgb or "").upper()
        return rgb.endswith("FFCDD2")
    except Exception:
        return False


def main():
    results = {}
    with sync_playwright() as pw:
        b = pw.chromium.connect_over_cdp("http://127.0.0.1:9222")
        ctx = b.contexts[0]
        for n in range(1, 13):
            fid = IDS[n]
            url = f"https://docs.google.com/spreadsheets/d/{fid}/export?format=xlsx"
            r = ctx.request.get(url)
            if r.status != 200:
                results[n] = {"error": f"HTTP {r.status}"}
                print(f"PC {n:02d}: GAGAL export HTTP {r.status}")
                continue
            wb = openpyxl.load_workbook(io.BytesIO(r.body()), data_only=True)
            names = wb.sheetnames
            has_val = "VALIDASI" in names
            info = {"tabs": names, "validasi": has_val}

            data_name = next((x for x in names if x not in NON_DATA), names[0])
            info["data_tab"] = data_name

            if has_val:
                vs = wb["VALIDASI"]
                issues = []
                for row in vs.iter_rows(min_row=2, max_row=min(vs.max_row, 60), values_only=True):
                    if row and len(row) >= 4 and (row[2] or row[3]):
                        if str(row[3]).strip() != "OK":
                            issues.append([row[0], row[1], row[2], row[3]])
                info["validasi_issues"] = issues[:8]
                info["validasi_rows"] = len(issues)

            ds = wb[data_name]
            notes_header = {}
            for c in range(1, 14):
                cell = ds.cell(1, c)
                if cell.comment:
                    notes_header[cell.coordinate] = cell.comment.text[:120]
            info["header_notes"] = notes_header
            info["note_i"] = notes_header.get("I1", "")

            red_cells, note_cells = [], []
            for r in range(2, min(ds.max_row, 60) + 1):
                for c in range(1, 14):
                    cell = ds.cell(r, c)
                    if is_red(cell):
                        red_cells.append(cell.coordinate)
                    if cell.comment:
                        note_cells.append((cell.coordinate, cell.comment.text[:70]))
            info["red_count"] = len(red_cells)
            info["red_sample"] = red_cells[:10]
            info["note_count"] = len(note_cells)
            info["note_sample"] = note_cells[:5]

            results[n] = info
            print(f"PC {n:02d}: VALIDASI={'ADA' if has_val else 'TIDAK'} | "
                  f"header notes={len(notes_header)} | sel merah={len(red_cells)} | "
                  f"sel note={len(note_cells)} | tab data='{data_name}' | note I='{info['note_i'][:40]}'")
            if has_val:
                print(f"        masalah terdeteksi: {info.get('validasi_rows', 0)} baris")
                for it in info.get("validasi_issues", [])[:3]:
                    print(f"        - baris {it[0]} ({it[1]}): {it[2]}")
            if info["note_sample"]:
                print(f"        contoh note: {info['note_sample'][0]}")
        b.close()

    (ROOT / "data" / "verify_install.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print("\nhasil lengkap -> data/verify_install.json")


if __name__ == "__main__":
    main()
