"""Gabungkan baris data dari 12 sumber ke output/master_karyawan.xlsx.

- Tab MASTER: 19 kolom (A..S), data A..O; formula P..S tidak disentuh.
- Tab Kumpulan Video: hanya baris yang Link Video YT (K) terisi dan link unik
  (kemunculan pertama dipertahankan).
- Header sumber (dengan spasi ekor) dipetakan ke kolom master.
- Tanggal dinormalisasi ke tipe datetime khusus kolom tanggal (D, L, M).
- Aggregator, genre, dan status dinormalkan ke nama baku dari
  config/validation_lists.json; ejaan lama Tuncore dipetakan ke TuneCore.
- Baris kosong dilewati.
- Baris MASTER ditulis berurutan (PC 01 dulu, lalu PC 02, dst) mulai baris 2.
"""
import json, re, sys
from datetime import datetime, timedelta
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
MFILE = ROOT / "output" / "master_karyawan.xlsx"
LISTS = json.loads((ROOT / "config" / "validation_lists.json").read_text(encoding="utf-8"))
AGGREGATOR_ALIASES = {"tuncore": "TuneCore"}

SRC_TO_COL = {
    "Worker ID": 3,                 # C
    "Update Date": 4,               # D
    "Song ID YT": 5,                # E
    "Judul Lagu Kompilasi": 6,      # F
    "Aggregator": 7,                # G
    "Genre": 8,                     # H
    "Nama Channel": 9,              # I
    "Link Channel YT": 10,          # J
    "Link Video YT": 11,            # K
    "Tanggal Video Dibuat": 12,     # L
    "Tanggal Video Dipublish": 13,  # M
    "Status Upload": 14,            # N
    "PIC channel": 15,              # O
}
DATE_COLS = {4, 12, 13}

REQUIRED_HEADERS = [
    "Worker ID", "Update Date", "Song ID YT", "Judul Lagu Kompilasi",
    "Aggregator", "Genre", "Nama Channel", "Link Channel YT", "Link Video YT",
    "Tanggal Video Dibuat", "Tanggal Video Dipublish", "Status Upload", "PIC channel",
]

# Kolom tab Kumpulan Video (9 kolom) -> indeks vals (0-based)
VIDEO_COLS = [0, 2, 4, 5, 8, 10, 13, 11, 12]

DATETIME_FORMATS = [
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
    "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
    "%Y/%m/%d", "%d.%m.%Y",
]


EXCEL_EPOCH = datetime(1899, 12, 30)
EXCEL_MAX_DAYS = 2958465  # 9999-12-31


def serial_to_dt(v):
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    if 1 <= n <= EXCEL_MAX_DAYS:
        return EXCEL_EPOCH + timedelta(days=n)
    return None


def norm_cell(v, is_date=False):
    """Trim teks; kolom tanggal diparse ke datetime. Return (value, issues)."""
    if v is None:
        return None, []
    if isinstance(v, bool):
        return v, []
    if isinstance(v, (int, float)):
        if is_date:
            dt = serial_to_dt(v)
            if dt is not None:
                return dt, []
            return None, [f"nilai tanggal tidak valid: {v}"]
        return v, []
    if isinstance(v, str):
        s = v.strip()
        if s == "":
            return None, []
        if is_date:
            for fmt in DATETIME_FORMATS:
                try:
                    return datetime.strptime(s, fmt), []
                except ValueError:
                    pass
            return s, [f"format tanggal tidak dikenali: '{s}'"]
        return s, []
    return v, []


def canon_value(value, allowed, aliases=None):
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    for item in allowed:
        if item.lower() == s.lower():
            return item
    if aliases and s.lower() in aliases:
        return aliases[s.lower()]
    return s


def canon_genres(value):
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    parts = [p.strip() for p in re.split(r"[,;/]+", s) if p.strip()]
    if not parts:
        return s
    return ", ".join(canon_value(p, LISTS["genre"]) for p in parts)


def is_blank_row(row):
    return all(v is None or (isinstance(v, str) and v.strip() == "") for v in row)


def is_data_row(row):
    """Baris dianggap data hanya jika punya Worker ID (A), Song ID (C), atau Judul (D)."""
    for idx in (0, 2, 3):
        v = row[idx] if idx < len(row) else None
        if v is not None and not (isinstance(v, str) and v.strip() == ""):
            return True
    return False


def preflight_header(header, pc):
    missing = [h for h in REQUIRED_HEADERS if h not in header]
    if "Link YT" in header:
        raise RuntimeError(
            f"PC {pc:02d}: header masih memakai 'Link YT' lama. "
            "Jalankan `python scripts/fetch_sheets.py` dulu agar sumber menjadi 13 kolom."
        )
    if missing:
        raise RuntimeError(
            f"PC {pc:02d}: header belum lengkap. Hilang: {', '.join(missing)}. "
            "Jalankan `python scripts/fetch_sheets.py` dulu."
        )


def clear_video_tab(video):
    last = max(video.max_row, 1000)
    for r in range(2, last + 1):
        for c in range(1, 10):
            video.cell(row=r, column=c).value = None


def main():
    wb_m = openpyxl.load_workbook(MFILE)
    if "MASTER" not in wb_m.sheetnames:
        raise RuntimeError("Tab MASTER tidak ada di master lokal.")
    if "Kumpulan Video" not in wb_m.sheetnames:
        raise RuntimeError("Tab Kumpulan Video tidak ada di master lokal; jalankan build_master.py dulu.")
    master = wb_m["MASTER"]
    video = wb_m["Kumpulan Video"]
    manifest = json.loads((ROOT / "data" / "manifest.json").read_text(encoding="utf-8"))

    # bersihkan nilai A..O lama (formula P..S jangan disentuh)
    last_old = master.max_row
    for r in range(2, last_old + 1):
        for c in range(1, 16):
            master.cell(row=r, column=c).value = None
    clear_video_tab(video)

    warn = []
    total = 0
    video_total = 0
    counts = {}
    cursor = 2
    video_cursor = 2
    seen_links = set()  # dedup global lintas PC

    for n in range(1, 13):
        src = RAW / f"src{n:02d}.xlsx"
        if not src.exists():
            warn.append(f"PC{n:02d}: file sumber tidak ada ({src.name})")
            counts[f"PC {n:02d}"] = 0
            continue

        wb_s = openpyxl.load_workbook(src, data_only=True)
        sh = wb_s.worksheets[0]
        rows = [list(r) for r in sh.iter_rows(values_only=True)]
        if not rows:
            warn.append(f"PC {n:02d}: file sumber kosong")
            counts[f"PC {n:02d}"] = 0
            continue
        header = [str(h).strip() if h else "" for h in rows[0]]
        preflight_header(header, n)

        # --- refresh tab RAW_PCxx ---
        rawname = f"RAW_PC{n:02d}"
        if rawname not in wb_m.sheetnames:
            raise RuntimeError(f"Tab {rawname} tidak ada di master lokal.")
        rawws = wb_m[rawname]
        if rawws.max_row > 1:
            rawws.delete_rows(2, rawws.max_row - 1)
        rr = 2
        skipped = 0
        for row in rows[1:]:
            if not is_data_row(row):
                if not is_blank_row(row):
                    skipped += 1
                continue
            for j, v in enumerate(row, 1):
                if v is not None and v != "":
                    rawws.cell(row=rr, column=j, value=v)
            rr += 1

        # --- tulis ke MASTER + Kumpulan Video ---
        n_rows = 0
        for row in rows[1:]:
            if not is_data_row(row):
                continue
            vals = [None] * 15
            vals[0] = f"PC {n:02d}"
            issues_row = []
            for j, hname in enumerate(header):
                if not hname:
                    continue
                col = SRC_TO_COL.get(hname)
                if col is None:
                    issues_row.append(f"kolom sumber '{hname}' tidak dipetakan")
                    continue
                v, iss = norm_cell(row[j] if j < len(row) else None,
                                   is_date=(col in DATE_COLS))
                vals[col - 1] = v
                if iss:
                    issues_row.extend(iss)
            vals[6] = canon_value(vals[6], LISTS["aggregator"], AGGREGATOR_ALIASES)
            vals[7] = canon_genres(vals[7])
            vals[13] = canon_value(vals[13], LISTS["status"])
            for j, v in enumerate(vals, 1):
                master.cell(row=cursor, column=j, value=v)

            # Kumpulan Video: hanya jika Link Video YT (K) terisi dan belum pernah muncul
            link_video = vals[10]
            if link_video is not None and str(link_video).strip() != "":
                if link_video not in seen_links:
                    seen_links.add(link_video)
                    for j, idx in enumerate(VIDEO_COLS, 1):
                        video.cell(row=video_cursor, column=j, value=vals[idx])
                    video_cursor += 1
                    video_total += 1

            if issues_row:
                warn.append(f"PC {n:02d} baris master {cursor}: " +
                            "; ".join(sorted(set(issues_row))))
            cursor += 1
            n_rows += 1
        if skipped:
            warn.append(f"PC {n:02d}: {skipped} baris tanpa Worker ID/Song ID/Judul diabaikan")
        counts[f"PC {n:02d}"] = n_rows
        total += n_rows

    wb_m.save(MFILE)
    print("gabungan data selesai ->", MFILE)
    print("total baris MASTER:", total)
    print("total baris Kumpulan Video:", video_total)
    for k, c in counts.items():
        print(f"  {k}: {c} baris")
    if warn:
        print("peringatan:")
        for w in warn:
            print("  -", w)
    return 0


if __name__ == "__main__":
    sys.exit(main())
